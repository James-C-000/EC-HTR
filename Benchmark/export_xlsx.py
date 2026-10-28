#!/usr/bin/env python3
"""
Excel deliverable from a production JSONL: one row per scan, taken from its latest JSONL row.

Only reads the JSONL, so it can run during production.py too. Rows are sorted by source_path.
Failed files are kept with their status and error (and no output), so none goes missing silently.
Columns added after the production fields:
  province_normalized  the raw province (2022 inventory name or 2014 abbreviation) as its Organized folder name
  run                  the JSONL the row came from, by file name (main, extras)
  scans_of_form        how many exported rows share this filename
  alternate_scan       True for rows from any JSONL after the first (pass main.jsonl first, then extras.jsonl)
Needs openpyxl (pip install openpyxl).

  python Benchmark/export_xlsx.py runs/main.jsonl                  # writes runs/main.xlsx
  python Benchmark/export_xlsx.py runs/main.jsonl runs/extras.jsonl --xlsx runs/all.xlsx
"""

import argparse
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from production import FIELDS, read_jsonl

MAX_CELL = 32_767   # Excel's limit on characters in a cell
EXTRA = ["province_normalized", "run", "scans_of_form", "alternate_scan"]

# Raw province -> Organized folder name (spaces for underscores). Built from the 26 distinct values in
# main.jsonl and extras.jsonl: 2022 inventory names, and 2014 abbreviations for stations only in 2014.
PROVINCE_2022 = {
    "ALBERTA": "Alberta", "BRITISH COLUMBIA": "British Columbia", "MANITOBA": "Manitoba",
    "NEW BRUNSWICK": "New Brunswick", "NEWFOUNDLAND": "Newfoundland and Labrador",
    "NORTHWEST TERRITORIES": "Northwest Territories", "NOVA SCOTIA": "Nova Scotia", "NUNAVUT": "Nunavut",
    "ONTARIO": "Ontario", "PRINCE EDWARD ISLAND": "Prince Edward Island", "QUEBEC": "Quebec",
    "SASKATCHEWAN": "Saskatchewan", "YUKON TERRITORY": "Yukon",
}
PROVINCE_2014 = {
    "ALTA": "Alberta", "BC": "British Columbia", "MAN": "Manitoba", "NB": "New Brunswick",
    "NFLD": "Newfoundland and Labrador", "NWT": "Northwest Territories", "NS": "Nova Scotia", "NU": "Nunavut",
    "ONT": "Ontario", "PEI": "Prince Edward Island", "QUE": "Quebec", "SASK": "Saskatchewan", "YT": "Yukon",
}
PROVINCE = {**PROVINCE_2022, **PROVINCE_2014}
LENGTH_NOTE = ("finish_reason=length rows (output hit the 512-token cap) are mostly repetition loops; their output "
               "is kept unedited. has_remarks is True for most of them, including loops of routine weather "
               "(e.g. 'Fair & warm') that the scope excludes. Filter on finish_reason when counting remarks.")


def cell(ws, key, v):
    if isinstance(v, list):                       # native_wh, sent_wh
        return " x ".join(map(str, v))
    if key in ("year", "month") and isinstance(v, str) and v.isdigit():
        return int(v)
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)[:MAX_CELL]
        if v.startswith("="):                     # a ditto mark is text, not a formula
            c = WriteOnlyCell(ws, v)
            c.data_type = "s"
            return c
    return v


def main():
    ap = argparse.ArgumentParser(description="Production JSONL -> Excel (latest row per source_path)")
    ap.add_argument("jsonl", type=Path, nargs="+", help="One or more production JSONL files")
    ap.add_argument("--xlsx", type=Path, help="Default: the first JSONL with .xlsx")
    args = ap.parse_args()
    xlsx = args.xlsx or args.jsonl[0].with_suffix(".xlsx")

    # source_path -> (file, byte offset, filename) of its latest row; offsets keep memory small on 570k rows
    latest = {}
    for i, path in enumerate(args.jsonl):
        for offset, rec in read_jsonl(path):
            latest[rec["source_path"]] = (i, offset, rec["filename"])
    scans = Counter(fn for _, _, fn in latest.values())
    runs = [p.stem for p in args.jsonl]

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Results")
    ws.freeze_panes = "A2"
    for i, key in enumerate(FIELDS + EXTRA, 1):
        ws.column_dimensions[get_column_letter(i)].width = {"source_path": 60, "output": 80, "error": 40,
                                                            "province_normalized": 26}.get(key, 14)
    ws.append(FIELDS + EXTRA)
    status, remarks, length = Counter(), Counter(), 0
    by_run, province_2014 = Counter(), 0
    files = [open(p, "rb") for p in args.jsonl]
    try:
        for sp in sorted(latest):
            i, offset, fn = latest[sp]
            files[i].seek(offset)
            rec = json.loads(files[i].readline())
            status[rec["status"]] += 1
            if rec["status"] == "ok":
                remarks["with remarks" if rec["has_remarks"] else "NO_REMARKS"] += 1
                length += rec["finish_reason"] == "length"
            if rec["province"] not in PROVINCE:
                raise ValueError(f"{sp}: unmapped province {rec['province']!r}; add it to PROVINCE")
            province_2014 += rec["province"] in PROVINCE_2014
            by_run[runs[i]] += 1
            ws.append([cell(ws, k, rec.get(k)) for k in FIELDS]
                      + [PROVINCE[rec["province"]], runs[i], scans[fn], i > 0])
    finally:
        for f in files:
            f.close()
    ws.auto_filter.ref = f"A1:{get_column_letter(len(FIELDS) + len(EXTRA))}{len(latest) + 1}"

    s = wb.create_sheet("Summary")
    s.column_dimensions["A"].width = 30
    s.column_dimensions["B"].width = 60
    s.append(["files", len(latest)])
    for k, v in sorted(status.items()):
        s.append([f"status {k}", v])
    for k, v in sorted(remarks.items()):
        s.append([f"ok, {k}", v])
    s.append(["ok, finish_reason=length", length])
    s.append(["source", ", ".join(str(p) for p in args.jsonl)])
    s.append(["exported (UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")])
    for r in runs:
        s.append([f"rows, run {r}", by_run[r]])
    multi = {fn: n for fn, n in scans.items() if n > 1}
    s.append(["forms with more than one scan", len(multi)])
    s.append(["rows in forms with more than one scan", sum(multi.values())])
    s.append(["rows with a 2014-style province", province_2014])
    note = WriteOnlyCell(s, LENGTH_NOTE)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    s.append(["note", note])
    wb.save(xlsx)

    print(f"{len(latest):,} files -> {xlsx}")
    print("  " + ", ".join(f"{k} {v:,}" for k, v in sorted(status.items())))
    print("  " + ", ".join(f"run {r} {by_run[r]:,}" for r in runs)
          + f"; {len(multi):,} forms with more than one scan ({sum(multi.values()):,} rows)"
          + f"; {province_2014:,} rows with a 2014-style province")
    not_done = sum(v for k, v in status.items() if k not in ("ok", "decode_error"))
    if not_done:
        print(f"WARNING: {not_done:,} files have no result yet (status other than ok/decode_error). "
              "Rerun production.py with the same arguments, then export again.")


if __name__ == "__main__":
    main()
