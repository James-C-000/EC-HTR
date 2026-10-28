#!/usr/bin/env python3
"""
Excel deliverable from the cleaned remarks: ec_remarks_clean.sqlite + ec_remarks_clean.parquet -> ec_remarks_clean.xlsx.

Inputs (read-only), from --in-dir (default: this folder), as written by clean_remarks.py in one run:
  ec_remarks_clean.sqlite   the rows: every column of its table remarks, sorted by source_path
  ec_remarks_clean.parquet  the `lines` column, for removed_lines; its source_path, output and
                            remarks_clean must match the SQLite row for row, or nothing is written

Output, in --out-dir (default: this folder), written under a temporary name and renamed when complete:
  ec_remarks_clean.xlsx
    Results  one row per scan (570,818): the remarks columns unchanged, with has_remarks,
             alternate_scan and has_remarks_clean as TRUE/FALSE (blank for failed files; an empty
             remarks_clean is a blank cell too, so use has_remarks_clean to tell them apart), plus
               removed_lines  every line the cleaning removed, in order, one per line as
                              "[rule] text" (the Excel form of `lines`; kept lines are in remarks_clean)
    Summary  counts, lines removed by rule, each input's size and SHA-256, the cleaning provenance,
             the export time and this script's git commit

Text is cleaned for Excel as in Production/export_xlsx.py: illegal characters are stripped, cells are
capped at 32,767 characters, and text that starts with "=" is stored as text. The Summary counts the
cells this changed.

Usage: python Post-Processing/export_clean_xlsx.py [--in-dir DIR] [--out-dir DIR]
Needs: duckdb, openpyxl
"""
import argparse
import hashlib
import os
import sqlite3
import subprocess
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import duckdb
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

MAX_CELL = 32_767   # Excel's limit on characters in a cell
BOOL_COLUMNS = {"has_remarks", "alternate_scan", "has_remarks_clean"}   # 0/1 in SQLite
WIDTHS = {"source_path": 60, "output": 80, "remarks_clean": 80, "removed_lines": 60, "error": 40,
          "province_normalized": 26}
NOTES = [
    ("remarks_clean", "Kept lines in their original order, one per distinct line; ' [×N]' marks a line that "
                      "occurred N times in a form that finished normally. Blank when nothing was kept "
                      "(has_remarks_clean FALSE) and for failed files (has_remarks_clean blank). "
                      "`output` is the unedited transcription."),
    ("removed_lines", "Every distinct line the cleaning removed, as '[rule] text'. Only junk is removed: "
                      "symbols (no letters or digits, e.g. ditto marks), null_entry (No, None, Nil), "
                      "cut_off_repeat (cut-off last line of a capped output that begins a kept line)."),
    ("finish_reason", "finish_reason=length rows (output hit the 512-token cap) are mostly repetition loops; "
                      "remarks_clean shows each of their lines once, without counts."),
]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while chunk := f.read(1 << 20):
            h.update(chunk)
    return h.hexdigest()


def script_commit():
    """HEAD of this script's repo, and whether the script differs from it (None, None outside git)."""
    here = Path(__file__).resolve()
    try:
        git = lambda *a: subprocess.run(["git", "-C", str(here.parent), *a], capture_output=True, text=True,
                                         check=True).stdout.strip()
        return git("rev-parse", "HEAD"), bool(git("status", "--porcelain", "--", here.name))
    except (OSError, subprocess.CalledProcessError):
        return None, None


class Cells:
    """Turns values into Excel cells, counting the text cells that had to change."""

    def __init__(self, ws):
        self.ws = ws
        self.stripped = Counter()
        self.truncated = Counter()

    def __call__(self, key, v):
        if key in BOOL_COLUMNS and v is not None:
            return bool(v)
        if not isinstance(v, str):
            return v
        s = ILLEGAL_CHARACTERS_RE.sub("", v)
        if s != v:
            self.stripped[key] += 1
        if len(s) > MAX_CELL:
            s = s[:MAX_CELL]
            self.truncated[key] += 1
        if s.startswith("="):                     # a ditto mark is text, not a formula
            c = WriteOnlyCell(self.ws, s)
            c.data_type = "s"
            return c
        return s


def removed_lines(lines):
    if lines is None:
        return None
    return "\n".join(f"[{d['removed_by']}] {d['text']}" for d in lines if d["removed_by"] is not None)


def batched(cur, size=5000):
    while rows := cur.fetchmany(size):
        yield from rows


def main():
    here = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser(description="Cleaned remarks (SQLite + Parquet) -> Excel (read-only on the inputs)")
    ap.add_argument("--in-dir", type=Path, default=here,
                    help="Folder with ec_remarks_clean.sqlite and ec_remarks_clean.parquet (default: %(default)s)")
    ap.add_argument("--out-dir", type=Path, default=here, help="Default: %(default)s")
    args = ap.parse_args()
    db_path, pq_path = args.in_dir / "ec_remarks_clean.sqlite", args.in_dir / "ec_remarks_clean.parquet"
    for p in (db_path, pq_path):
        if not p.is_file():
            sys.exit(f"missing input: {p}")
    args.out_dir.mkdir(parents=True, exist_ok=True)
    xlsx = args.out_dir / "ec_remarks_clean.xlsx"
    tmp = xlsx.with_name(f".{xlsx.name}.tmp")
    tmp.unlink(missing_ok=True)
    try:
        inputs = [{"path": str(p.resolve()), "size": p.stat().st_size, "sha256": sha256(p)} for p in (db_path, pq_path)]
        n = export(db_path, pq_path, tmp, inputs)
        for info in inputs:
            if Path(info["path"]).stat().st_size != info["size"]:
                raise RuntimeError(f"{info['path']} changed during the run; rerun once it is no longer being written")
        os.replace(tmp, xlsx)
    except BaseException:
        tmp.unlink(missing_ok=True)
        raise
    print(f"{n:,} rows -> {xlsx}")


def export(db_path: Path, pq_path: Path, dst: Path, inputs: list) -> int:
    db = sqlite3.connect(f"{db_path.resolve().as_uri()}?mode=ro", uri=True)
    con = duckdb.connect()
    try:
        cur = db.execute("SELECT * FROM remarks ORDER BY source_path")
        columns = [d[0] for d in cur.description]
        rules = [c.removeprefix("n_removed_") for c in columns if c.startswith("n_removed_")]   # as clean_remarks.py wrote them
        if "remarks_clean" not in columns or not rules:
            raise ValueError(f"{db_path.name}: the remarks table has no cleaning columns; is this clean_remarks.py output?")
        header = columns + ["removed_lines"]
        i_sp, i_out, i_clean = (columns.index(k) for k in ("source_path", "output", "remarks_clean"))
        pq = batched(con.execute("SELECT source_path, output, remarks_clean, lines FROM read_parquet(?) "
                                 "ORDER BY source_path", [str(pq_path)]))

        wb = Workbook(write_only=True)
        ws = wb.create_sheet("Results")
        ws.freeze_panes = "A2"
        for i, key in enumerate(header, 1):
            ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(key, 14)
        ws.append(header)
        cells = Cells(ws)
        status, has, has_clean, removed = Counter(), Counter(), Counter(), Counter()
        n = n_raw = n_kept = 0
        for row in batched(cur):
            p = next(pq, None)
            if p is None or p[:3] != (row[i_sp], row[i_out], row[i_clean]):
                raise ValueError(f"{db_path.name} and {pq_path.name} differ at {row[i_sp]}; "
                                 "use the two files written by the same clean_remarks.py run")
            rec = dict(zip(columns, row))
            status[rec["status"]] += 1
            has[rec["has_remarks"]] += 1
            has_clean[rec["has_remarks_clean"]] += 1
            n_raw += rec["n_lines_raw"]
            n_kept += rec["n_lines_clean"]
            for k in rules:
                removed[k] += rec[f"n_removed_{k}"]
            ws.append([cells(k, v) for k, v in zip(columns, row)]
                      + [cells("removed_lines", removed_lines(p[3]))])
            n += 1
        if next(pq, None) is not None:
            raise ValueError(f"{pq_path.name} has more rows than {db_path.name}")
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{n + 1}"

        s = wb.create_sheet("Summary")
        s.column_dimensions["A"].width = 34
        s.column_dimensions["B"].width = 90
        label = {True: "True", False: "False", None: "null (failed file)"}
        s.append(["files", n])
        for k, v in sorted(status.items()):
            s.append([f"status {k}", v])
        for k in (True, False, None):
            s.append([f"has_remarks {label[k]}", has[k]])
        for k in (True, False, None):
            s.append([f"has_remarks_clean {label[k]}", has_clean[k]])
        s.append(["output lines (non-empty)", n_raw])
        s.append(["distinct lines kept", n_kept])
        for k in rules:
            s.append([f"distinct lines removed: {k}", removed[k]])
        for k, v in sorted(cells.stripped.items()):
            s.append([f"Excel: {k} cells with illegal characters stripped", v])
        for k, v in sorted(cells.truncated.items()):
            s.append([f"Excel: {k} cells cut to {MAX_CELL:,} characters", v])
        if not cells.stripped and not cells.truncated:
            s.append(["Excel: cells changed (stripped or cut)", 0])
        for k, note in NOTES:
            c = WriteOnlyCell(s, note)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            s.append([f"note: {k}", c])
        for d in inputs:
            s.append(["source", d["path"]])
            s.append(["  size, sha256", f"{d['size']} bytes, {d['sha256']}"])
        for path, size, digest, cleaned, commit, modified in db.execute(
                "SELECT path, size, sha256, cleaned_utc, script_commit, script_modified FROM clean_provenance"):
            s.append(["cleaned from", path])
            s.append(["  size, sha256", f"{size} bytes, {digest}"])
            s.append(["  cleaned (UTC), clean_remarks.py commit",
                      f"{cleaned}, {commit}{' (modified)' if modified else ''}"])
        commit, modified = script_commit()
        s.append(["exported (UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")])
        s.append(["export_clean_xlsx.py commit", f"{commit}{' (modified)' if modified else ''}"])
        wb.save(dst)
    finally:
        db.close()
        con.close()

    print("  " + ", ".join(f"{k} {v:,}" for k, v in sorted(status.items()))
          + f"; has_remarks_clean True {has_clean[True]:,}")
    print("  lines: " + f"{n_raw:,} raw, {n_kept:,} distinct kept; removed "
          + ", ".join(f"{k} {removed[k]:,}" for k in rules))
    changed = sum(cells.stripped.values()) + sum(cells.truncated.values())
    if changed:
        print(f"  Excel: {sum(cells.stripped.values()):,} cells had illegal characters stripped, "
              f"{sum(cells.truncated.values()):,} were cut to {MAX_CELL:,} characters")
    return n


if __name__ == "__main__":
    main()
